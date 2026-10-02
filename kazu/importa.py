"""
Import kazu husi Excel (Admin deit).

	1. Admin download template (aba "Kazu" + aba lista kódigu)
	2. Upload → lee → valida kada liña → pratinjau (KazuImport RASCUNHO, la kria kazu)
	3. Konfirma → kria kazu ho status "Hein Verifikasaun" (fluxu hanesan: Admin verifika, Superadmin aprova)

Koordenada mamuk ka la kompletu → fatin aproksimadu (sentru suku, ka munisípiu) + marka gps_aproksimadu;
Admin hadia iha mapa (marka bele dada). Konsentimentu mamuk = LAE → kazu la tama portal públiku.
"""
import hashlib
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from config.rbac import SUPERADMIN
from custom.models import Munisipiu, PostuAdministrativu, Suku, Aldeia, TipuKonflitu, TipuRai
from kazu.forms import TL_LAT, TL_LON
from kazu.models import Kazu, KazuHistoria, KazuImport, UmaKainAfetada, SYNCED, VERIFIED
from kazu.services import generate_kode, _historia
from notification.utils import kirim_notif_role

MAX_LINA = 500
MAX_BYTES = 5 * 1024 * 1024
SIN = {'SIN', 'S', 'YES', 'Y', 'SIM', 'YA', '1', 'TRUE', 'LOOS'}

# (kolun, deskrisaun ba template, obrigatóriu)
KOLUN = [
	('data_relatoriu', 'Data relatóriu (AAAA-MM-LL)', True),
	('munisipiu', 'Munisípiu (kódigu ka naran)', True),
	('postu', 'Postu (kódigu ka naran)', False),
	('suku', 'Suku (kódigu ka naran)', False),
	('aldeia', 'Aldeia (kódigu ka naran)', False),
	('latitude', 'Latitude (ex. -8.5586)', False),
	('longitude', 'Longitude (ex. 125.5736)', False),
	('data_akontesimentu', 'Data akontesimentu (AAAA-MM-LL)', False),
	('tipu_konflitu', 'Tipu konflitu (kódigu, separa ho ;)', True),
	('tipu_seluk', 'Tipu "Seluk": esplika', False),
	('tipu_rai', 'Tipu rai (kódigu)', False),
	('deskrisaun', 'Deskrisaun insidente', True),
	('uma_kain', 'Uma-kain afetada', False),
	('mane', 'Mane', False),
	('feto', 'Feto', False),
	('labarik', 'Labarik', False),
	('konsentimentu', 'Konsentimentu (SIN / LAE)', False),
	('la_publika', 'La publika iha portal (SIN / LAE)', False),
	('titulu', 'Títulu (opsionál)', False),
	('observasaun', 'Observasaun', False),
]
NARAN_KOLUN = [k[0] for k in KOLUN]


# ══════════════ 1. Template ══════════════
def template_excel():
	wb = Workbook()
	ws = wb.active
	ws.title = 'Kazu'
	kab = PatternFill('solid', fgColor='0B4F3C')
	obrig = PatternFill('solid', fgColor='0E8A68')
	ws.append(NARAN_KOLUN)
	ws.append([d for _k, d, _o in KOLUN])
	for i, (_k, _d, o) in enumerate(KOLUN, start=1):
		c = ws.cell(row=1, column=i)
		c.font, c.fill = Font(bold=True, color='FFFFFF'), (obrig if o else kab)
		ws.cell(row=2, column=i).font = Font(italic=True, color='5E6B66')
		ws.column_dimensions[c.column_letter].width = 22
	ws.append(['2026-09-15', 'LIQ', '', '', '', '', '', '2026-09-10', 'DESLOKAMENTU', '', '', 'Ezemplu: família 12 hetan avizu atu sai husi rai.', 12, 30, 28, 20, 'SIN', 'LAE', '', 'Hamoos liña ezemplu ne\'e'])
	for c in ws[3]:
		c.font = Font(color='9AA59F')
	ws.freeze_panes = 'A3'

	# Aba lista (kódigu + naran) — ajuda prenxe
	def aba(naran, kab_list, rows):
		w = wb.create_sheet(naran)
		w.append(kab_list)
		for c in w[1]:
			c.font, c.fill = Font(bold=True, color='FFFFFF'), kab
		for r in rows:
			w.append(list(r))
		for col in w.columns:
			w.column_dimensions[col[0].column_letter].width = 28
		return w
	aba('Munisipiu', ['code', 'naran'], Munisipiu.active.values_list('code', 'name'))
	aba('Postu', ['code', 'naran', 'munisipiu'], PostuAdministrativu.active.values_list('code', 'name', 'munisipiu__code'))
	aba('Suku', ['code', 'naran', 'postu'], Suku.active.values_list('code', 'name', 'postu__code'))
	aba('TipuKonflitu', ['code', 'naran'], TipuKonflitu.active.values_list('code', 'name'))
	aba('TipuRai', ['code', 'naran'], TipuRai.active.values_list('code', 'name'))

	# Dropdown iha aba Kazu
	n_mun = Munisipiu.active.count() + 1
	for col, formula in (('B', f'=Munisipiu!$A$2:$A${n_mun}'), ('Q', '"SIN,LAE"'), ('R', '"SIN,LAE"')):
		dv = DataValidation(type='list', formula1=formula, allow_blank=True)
		ws.add_data_validation(dv)
		dv.add(f'{col}3:{col}{MAX_LINA + 2}')
	buf = BytesIO()
	wb.save(buf)
	return buf.getvalue()


# ══════════════ 2. Lee no valida ══════════════
def lee_excel(f):
	# Retorna (naran_file, sha256, liña sira [{kolun: valór}]) ka ValidationError
	if f.size > MAX_BYTES:
		raise ValidationError(_('File boot liu (máximu 5 MB).'))
	raw = f.read()
	if not f.name.lower().endswith('.xlsx') or raw[:2] != b'PK':
		raise ValidationError(_('File tenke Excel .xlsx (uza template).'))
	try:
		wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
	except Exception:
		raise ValidationError(_('Labele lee file Excel ne\'e.'))
	ws = wb['Kazu'] if 'Kazu' in wb.sheetnames else wb.worksheets[0]
	rows = list(ws.iter_rows(values_only=True))
	if not rows:
		raise ValidationError(_('File mamuk.'))
	kab = [str(c or '').strip().lower() for c in rows[0]]
	if 'munisipiu' not in kab or 'deskrisaun' not in kab:
		raise ValidationError(_('Kolun la tuir template (presiza "munisipiu", "deskrisaun"…). Download template.'))
	lina = []
	for n, r in enumerate(rows[1:], start=2):
		d = {k: v for k, v in zip(kab, r) if k in NARAN_KOLUN}
		if not any(v not in (None, '') for v in d.values()):
			continue
		if n == 2 and str(d.get('data_relatoriu') or '').startswith('Data relat'):
			continue                                    # liña esplikasaun template
		d['_lina'] = n
		lina.append(d)
	if len(lina) > MAX_LINA:
		raise ValidationError(_('Liña barak liu (máximu %(n)s). Fahe ba file barak.') % {'n': MAX_LINA})
	return f.name[:255], hashlib.sha256(raw).hexdigest(), lina


def _testu(v):
	return '' if v is None else str(v).strip()


def _data(v):
	if v in (None, ''):
		return None
	if isinstance(v, datetime):
		return v.date()
	if isinstance(v, date):
		return v
	t = _testu(v)
	for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
		try:
			return datetime.strptime(t, fmt).date()
		except ValueError:
			pass
	raise ValueError(t)


def _int(v):
	if v in (None, ''):
		return None
	n = int(float(str(v).replace(',', '.')))
	if n < 0:
		raise ValueError(v)
	return n


def _dec(v):
	if v in (None, ''):
		return None
	return Decimal(str(v).replace(',', '.').strip()).quantize(Decimal('0.000001'))


class _Buka:
	# Buka objetu husi kódigu ka naran (la sensível ba letra boot/kiik)
	def __init__(self, qs):
		self.code = {o.code.upper(): o for o in qs}
		self.naran = {}
		for o in qs:
			self.naran.setdefault(o.name.strip().upper(), []).append(o)

	def __call__(self, v, filtru=None):
		t = _testu(v).upper()
		if not t:
			return None
		o = self.code.get(t)
		if o:
			return o
		opsaun = [x for x in self.naran.get(t, []) if not filtru or filtru(x)]
		return opsaun[0] if len(opsaun) == 1 else False   # False = la hetan ka naran duplikadu


def valida(lina):
	# Kada liña: {'lina', 'ok', 'erru': [...], 'avizu': [...], 'dadus': {... id/valór ba kria}, 'rezumu': {...}}
	buka_mun = _Buka(list(Munisipiu.active.all()))
	buka_postu = _Buka(list(PostuAdministrativu.active.select_related('munisipiu')))
	buka_suku = _Buka(list(Suku.active.select_related('postu')))
	buka_aldeia = _Buka(list(Aldeia.active.select_related('suku')))
	buka_tipu = _Buka(list(TipuKonflitu.active.all()))
	buka_rai = _Buka(list(TipuRai.active.all()))
	hoje = timezone.localdate()
	haree = set()
	rezultadu = []
	for d in lina:
		erru, avizu, out = [], [], {}

		# Data
		for campo, label in (('data_relatoriu', _('Data relatóriu')), ('data_akontesimentu', _('Data akontesimentu'))):
			try:
				out[campo] = _data(d.get(campo))
			except ValueError:
				erru.append(_('%(c)s la validu (uza AAAA-MM-LL).') % {'c': label})
				out[campo] = None
		if not out['data_relatoriu'] and not any(_('Data relatóriu') in e for e in erru):
			erru.append(_('Data relatóriu obrigatóriu.'))
		if out['data_relatoriu'] and out['data_relatoriu'] > hoje:
			erru.append(_('Data relatóriu labele iha futuru.'))
		if out['data_relatoriu'] and out['data_akontesimentu'] and out['data_akontesimentu'] > out['data_relatoriu']:
			erru.append(_('Data akontesimentu labele liu data relatóriu.'))

		# Lokalizasaun (tenke tuir malu)
		mun = buka_mun(d.get('munisipiu'))
		if not mun:
			erru.append(_('Munisípiu "%(v)s" la hetan.') % {'v': _testu(d.get('munisipiu'))} if _testu(d.get('munisipiu')) else _('Munisípiu obrigatóriu.'))
		postu = buka_postu(d.get('postu'), lambda o: mun and o.munisipiu_id == mun.pk)
		suku = buka_suku(d.get('suku'), lambda o: not postu or o.postu_id == postu.pk)
		aldeia = buka_aldeia(d.get('aldeia'), lambda o: not suku or o.suku_id == suku.pk)
		for obj, campo, label in ((postu, 'postu', _('Postu')), (suku, 'suku', _('Suku')), (aldeia, 'aldeia', _('Aldeia'))):
			if obj is False:
				erru.append(_('%(c)s "%(v)s" la hetan (ka naran duplikadu: uza kódigu).') % {'c': label, 'v': _testu(d.get(campo))})
		if mun and postu and postu.munisipiu_id != mun.pk:
			erru.append(_('Postu la iha munisípiu ne\'e.'))
		if postu and suku and suku.postu_id != postu.pk:
			erru.append(_('Suku la iha postu ne\'e.'))
		if suku and not postu:
			postu = suku.postu                              # postu husi suku
		if suku and aldeia and aldeia.suku_id != suku.pk:
			erru.append(_('Aldeia la iha suku ne\'e.'))

		# GPS: kompletu no iha Timor-Leste; se lae → fatin aproksimadu
		lat = lon = None
		try:
			lat, lon = _dec(d.get('latitude')), _dec(d.get('longitude'))
		except (InvalidOperation, ValueError):
			avizu.append(_('Koordenada la validu — uza fatin aproksimadu.'))
		if (lat is None) != (lon is None):
			avizu.append(_('Koordenada la kompletu — uza fatin aproksimadu.'))
			lat = lon = None
		if lat is not None and not (TL_LAT[0] <= lat <= TL_LAT[1] and TL_LON[0] <= lon <= TL_LON[1]):
			avizu.append(_('Koordenada li\'ur Timor-Leste — uza fatin aproksimadu.'))
			lat = lon = None
		aproksimadu = lat is None
		if aproksimadu and mun:
			sentru = suku if suku and suku.latitude is not None else mun
			if sentru.latitude is None:
				erru.append(_('Koordenada mamuk no munisípiu la iha sentru: hatama latitude/longitude.'))
			else:
				lat, lon = sentru.latitude, sentru.longitude
				if not any('aproksimadu' in a for a in avizu):
					avizu.append(_('Koordenada mamuk — uza sentru %(f)s. Hadia iha mapa hafoin import.') % {'f': sentru.name})

		# Tipu konflitu (kódigu barak, separa ho ; ka ,)
		tipu = []
		for t in [x for x in _testu(d.get('tipu_konflitu')).replace(',', ';').split(';') if x.strip()]:
			o = buka_tipu(t)
			if o:
				tipu.append(o)
			else:
				erru.append(_('Tipu konflitu "%(v)s" la hetan.') % {'v': t.strip()})
		if not tipu and not any(_('Tipu konflitu') in e for e in erru):
			erru.append(_('Tipu konflitu obrigatóriu.'))
		if any(t.presiza_esplika for t in tipu) and not _testu(d.get('tipu_seluk')):
			erru.append(_('Favor esplika tipu konflitu "Seluk".'))
		rai = buka_rai(d.get('tipu_rai'))
		if rai is False:
			erru.append(_('Tipu rai "%(v)s" la hetan.') % {'v': _testu(d.get('tipu_rai'))})

		deskrisaun = _testu(d.get('deskrisaun'))
		if not deskrisaun:
			erru.append(_('Deskrisaun obrigatóriu.'))

		# Populasaun
		pop = {}
		for campo in ('uma_kain', 'mane', 'feto', 'labarik'):
			try:
				pop[campo] = _int(d.get(campo))
			except (ValueError, TypeError):
				erru.append(_('"%(c)s" tenke númeru pozitivu.') % {'c': campo})
				pop[campo] = None

		konsentimentu = _testu(d.get('konsentimentu')).upper() in SIN
		if not konsentimentu:
			avizu.append(_('La iha konsentimentu — kazu la sei mosu iha portal públiku.'))

		# Duplikadu (iha file ka iha sistema)
		chave = (mun.pk if mun else None, out['data_relatoriu'], deskrisaun[:80].upper())
		if chave in haree:
			avizu.append(_('Liña hanesan iha file ne\'e (duplikadu?).'))
		haree.add(chave)
		if mun and out['data_relatoriu'] and Kazu.objects.filter(munisipiu=mun, data_relatoriu=out['data_relatoriu'], deskrisaun__iexact=deskrisaun).exists():
			avizu.append(_('Kazu hanesan iha sistema ona (duplikadu?).'))

		rezultadu.append({
			'lina': d['_lina'], 'ok': not erru, 'erru': erru, 'avizu': avizu,
			'dadus': {
				'titulu': _testu(d.get('titulu'))[:200],
				'data_relatoriu': out['data_relatoriu'].isoformat() if out['data_relatoriu'] else None,
				'data_akontesimentu': out['data_akontesimentu'].isoformat() if out['data_akontesimentu'] else None,
				'munisipiu': mun.pk if mun else None, 'postu': postu.pk if postu else None,
				'suku': suku.pk if suku else None, 'aldeia': aldeia.pk if aldeia else None,
				'latitude': str(lat) if lat is not None else None, 'longitude': str(lon) if lon is not None else None,
				'gps_aproksimadu': aproksimadu,
				'tipu_konflitu': [t.pk for t in tipu], 'tipu_seluk': _testu(d.get('tipu_seluk'))[:255],
				'tipu_rai': rai.pk if rai else None, 'deskrisaun': deskrisaun,
				'konsentimentu': konsentimentu, 'la_publika': _testu(d.get('la_publika')).upper() in SIN,
				'observasaun': _testu(d.get('observasaun')),
				**pop,
			},
			'rezumu': {
				'munisipiu': mun.name if mun else _testu(d.get('munisipiu')),
				'suku': suku.name if suku else '',
				'tipu': ', '.join(t.name for t in tipu),
				'deskrisaun': deskrisaun[:90],
			},
		})
	return rezultadu


def kria_pratinjau(f, user):
	naran, sha, lina = lee_excel(f)
	rez = valida(lina)
	return KazuImport.objects.create(
		naran_file=naran, sha256=sha, dadus=rez, created_by=user, total=len(rez),
		total_ok=sum(1 for r in rez if r['ok']),
		total_aproksimadu=sum(1 for r in rez if r['ok'] and r['dadus']['gps_aproksimadu']))


# ══════════════ 3. Konfirma ══════════════
@transaction.atomic
def konfirma(imp, user):
	imp = KazuImport.objects.select_for_update().get(pk=imp.pk)
	if imp.status != KazuImport.RASCUNHO:
		raise ValidationError(_('Import ne\'e konfirma ka kansela ona.'))
	kria = 0
	for r in imp.dadus:
		if not r['ok']:
			continue
		d = r['dadus']
		kazu = Kazu(
			titulu=d['titulu'], status=SYNCED, synced_at=timezone.now(),
			data_relatoriu=date.fromisoformat(d['data_relatoriu']),
			data_akontesimentu=date.fromisoformat(d['data_akontesimentu']) if d['data_akontesimentu'] else None,
			munisipiu_id=d['munisipiu'], postu_id=d['postu'], suku_id=d['suku'], aldeia_id=d['aldeia'],
			latitude=Decimal(d['latitude']), longitude=Decimal(d['longitude']), gps_aproksimadu=d['gps_aproksimadu'],
			tipu_seluk=d['tipu_seluk'], tipu_rai_id=d['tipu_rai'], deskrisaun=d['deskrisaun'],
			konsentimentu=d['konsentimentu'], la_publika=d['la_publika'], observasaun=d['observasaun'],
			created_by=user, importasaun=imp)
		kazu.kode = generate_kode(kazu)
		kazu.save()
		kazu.tipu_konflitu.set(d['tipu_konflitu'])
		if any(d.get(c) is not None for c in ('uma_kain', 'mane', 'feto', 'labarik')):
			total = (d.get('mane') or 0) + (d.get('feto') or 0)
			UmaKainAfetada.objects.create(kazu=kazu, uma_kain=d.get('uma_kain'), mane=d.get('mane'), feto=d.get('feto'),
				labarik=d.get('labarik'), total_ema=total or None)
		_historia(kazu, user, '', SYNCED, _('Import Excel: %(f)s · liña %(n)s') % {'f': imp.naran_file, 'n': r['lina']})
		kria += 1
	imp.status = KazuImport.REMATA
	imp.konfirma_iha = timezone.now()
	imp.save(update_fields=['status', 'konfirma_iha'])
	if kria:
		msg = lambda: _('Admin import kazu %(n)s husi Excel (%(f)s). Hein verifikasaun no aprovasaun.') % {'n': kria, 'f': imp.naran_file}
		transaction.on_commit(lambda: kirim_notif_role([SUPERADMIN], 'KAZU_FOUN', msg, url='/kazu/', actor=user, email=False))
	return kria


# ══════════════ 4. Hadia lokasaun (mapa) ══════════════
def pode_muda_lokasaun(kazu):
	# Admin bele muda fatin molok kazu aprova (depois aprova, dadus fiksu)
	return kazu.status in (SYNCED, VERIFIED)


@transaction.atomic
def muda_lokasaun(kazu, user, lat, lon):
	if not (TL_LAT[0] <= lat <= TL_LAT[1] and TL_LON[0] <= lon <= TL_LON[1]):
		raise ValidationError(_('Koordenada GPS la iha territóriu Timor-Leste.'))
	antes = f'{kazu.latitude}, {kazu.longitude}'
	kazu.latitude, kazu.longitude, kazu.gps_aproksimadu = lat, lon, False
	kazu.save(update_fields=['latitude', 'longitude', 'gps_aproksimadu'])
	_historia(kazu, user, kazu.status, kazu.status, _('Lokasaun muda: %(a)s → %(b)s, %(c)s') % {'a': antes, 'b': lat, 'c': lon}, tipu=KazuHistoria.EDITA)
	return kazu
