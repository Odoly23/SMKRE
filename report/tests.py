from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from config.rbac import ADMIN, SUPERADMIN, INVESTIGADOR, ANALISTA, OFISIAL_LEGAL
from config.testing import setup_master, make_user
from custom.models import Munisipiu, TipuKonflitu, TipuEviksaun
from kazu.models import Kazu, UmaKainAfetada, InsidenteEviksaun


class ReportBase(TestCase):
	def setUp(self):
		setup_master()
		self.inv = make_user(INVESTIGADOR)
		self.analista = make_user(ANALISTA)
		self.dil = Munisipiu.objects.get(code='DIL')
		self.liq = Munisipiu.objects.get(code='LIQ')
		self.deslok = TipuKonflitu.objects.get(code='GOV-RAI-ESTADU')        # tipu eviksaun
		self.tradisional = TipuKonflitu.objects.get(code='KOM-KOM')
		hoje = timezone.localdate()
		# 3 kazu Dili, 1 Liquiçá, 1 rascunho (la tama), 1 kanseladu (la tama)
		for i, (mun, status, tipu, urj) in enumerate([
				(self.dil, 'SYNCED', self.deslok, True), (self.dil, 'VERIFIED', self.tradisional, False),
				(self.dil, 'COMPLETED', self.deslok, False), (self.liq, 'APPROVED', self.tradisional, False),
				(self.liq, 'DRAFT', self.deslok, False), (self.liq, 'CANCELED', self.deslok, False)]):
			k = Kazu.objects.create(munisipiu=mun, status=status, created_by=self.inv, kode=f'T-{i}', urjente=urj,
				data_relatoriu=hoje, latitude=-8.55, longitude=125.57, konsentimentu=True)
			k.tipu_konflitu.set([tipu])
			UmaKainAfetada.objects.create(kazu=k, uma_kain=10, total_ema=40, mane=20, feto=20, labarik=15)
		InsidenteEviksaun.objects.create(kazu=Kazu.objects.get(kode='T-3'), tipu_eviksaun=TipuEviksaun.objects.first())


class DashTest(ReportBase):
	def test_investigador_labele_asesu(self):
		self.client.force_login(self.inv)
		for name in ('report-dash', 'report-chart', 'report-mapa', 'report-list', 'report-export-excel'):
			self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)
		self.assertEqual(self.client.get('/api/report/kazu/mun/').status_code, 403)

	def test_papel_dashboard_bele_asesu(self):
		for role in (SUPERADMIN, ADMIN, ANALISTA, OFISIAL_LEGAL):
			self.client.force_login(make_user(role, email=f'{role}.dash@teste.tl'))
			self.assertEqual(self.client.get(reverse('report-dash')).status_code, 200)

	def test_kontador(self):
		self.client.force_login(self.analista)
		r = self.client.get(reverse('report-dash'))
		self.assertEqual(r.context['total'], 4)                 # la konta rascunho no kanseladu
		self.assertEqual(r.context['tot_uma_kain'], 40)
		self.assertEqual(r.context['tot_remata'], 1)
		self.assertEqual(r.context['tot_urjente'], 1)
		self.assertEqual(r.context['tot_aktivu'], 2)             # SYNCED deslok + APPROVED ho insidente
		mun = {m.code: n for m, n in r.context['objects3']}
		self.assertEqual(mun['DIL'], 3)
		self.assertEqual(mun['LIQ'], 1)

	def test_filtru_munisipiu(self):
		self.client.force_login(self.analista)
		r = self.client.get(reverse('report-dash'), {'munisipiu': self.liq.pk})
		self.assertEqual(r.context['total'], 1)

	def test_lista_klik_numeru(self):
		self.client.force_login(self.analista)
		r = self.client.get(reverse('report-mun-list', args=[self.dil.pk]))
		self.assertEqual(len(r.context['objects']), 3)
		r = self.client.get(reverse('report-status-list', args=['COMPLETED']))
		self.assertEqual(len(r.context['objects']), 1)
		r = self.client.get(reverse('report-tipu-list', args=[self.deslok.pk]))
		self.assertEqual(len(r.context['objects']), 2)


class APITest(ReportBase):
	def test_api_mun_no_mapa(self):
		self.client.force_login(self.analista)
		d = self.client.get('/api/report/kazu/mun/').json()
		self.assertEqual(dict(zip(d['label'], d['obj']))['Dili'], 3)
		m = self.client.get('/api/report/kazu/mapa/').json()
		total = {x['code']: x['total'] for x in m['munisipiu']}
		self.assertEqual(total['DIL'], 3)                        # regresaun: order_by() iha GROUP BY
		self.assertEqual(len(m['pins']), 4)

	def test_api_tendensia_no_afetadu(self):
		self.client.force_login(self.analista)
		t = self.client.get('/api/report/kazu/tendensia/').json()
		self.assertEqual(len(t['label']), 12)
		self.assertEqual(t['total'][-1], 4)
		a = self.client.get('/api/report/kazu/afetadu/').json()
		self.assertEqual(dict(zip(a['label'], a['mane']))['Dili'], 60)


class ExportTest(ReportBase):
	def test_excel(self):
		self.client.force_login(self.analista)
		r = self.client.get(reverse('report-export-excel'), {'munisipiu': self.dil.pk})
		self.assertEqual(r.status_code, 200)
		self.assertIn('spreadsheetml', r['Content-Type'])
		ws = load_workbook(BytesIO(r.content)).active
		self.assertEqual(ws.max_row, 4)                          # header + 3 kazu Dili
		self.assertEqual(ws['A2'].value[:2], 'T-')

	def test_excel_kolun_rbr_sein_kontaktu(self):
		from kazu.models import KazuSegimentu
		k = Kazu.objects.filter(munisipiu=self.dil).exclude(status__in=['DRAFT', 'CANCELED']).first()
		Kazu.objects.filter(pk=k.pk).update(eviksaun_akontese='SIN', portavos_kontaktu='77009988')
		KazuSegimentu.objects.create(kazu=k, status_rezolusaun='SEIDAUK', fokal_telemovel='77112233')
		self.client.force_login(self.analista)
		ws = load_workbook(BytesIO(self.client.get(reverse('report-export-excel'), {'munisipiu': self.dil.pk}).content)).active
		kab = [c.value for c in ws[1]]
		self.assertEqual(kab[:3], ['Identifikasaun kazu', 'Kazu nia numero', 'Numero Referensia Kazu (auto)'])   # kolun Airtable
		self.assertEqual(kab.count('Status Kazu'), 1)
		valores = [str(c.value) for row in ws.iter_rows(min_row=2) for c in row]
		self.assertIn('Seidauk resolve', valores)
		self.assertNotIn('77009988', valores)                      # kontaktu ema la sai
		self.assertNotIn('77112233', valores)
		# Admin haree kontaktu (formatu kompletu ba import fali)
		self.client.force_login(make_user(ADMIN, email='adm2@teste.tl'))
		ws = load_workbook(BytesIO(self.client.get(reverse('report-export-excel'), {'munisipiu': self.dil.pk}).content)).active
		self.assertIn('77112233', [str(c.value) for row in ws.iter_rows(min_row=2) for c in row])


class MapaTest(ReportBase):
	def test_mapa_pin_no_hotspot(self):
		self.client.force_login(self.analista)
		r = self.client.get(reverse('report-mapa'))
		self.assertEqual(r.status_code, 200)
		self.assertEqual(len(r.context['mapobjects']), 4)                # rascunho no kanseladu la tama
		self.assertEqual(dict((m.code, n) for m, n in r.context['munobjects'])['DIL'], 3)
		self.assertContains(r, 'L.control.layers')

	def test_popup_escape_anti_xss(self):
		Kazu.objects.filter(kode='T-0').update(kode='<img src=x onerror=alert(1)>')
		self.client.force_login(self.analista)
		r = self.client.get(reverse('report-mapa'))
		self.assertNotIn(b'<img src=x onerror', r.content)
		self.assertIn(b'\\u0026lt\\u003Bimg', r.content)


class LiveTest(TestCase):
	def setUp(self):
		from django.core.cache import cache
		cache.clear()
		setup_master()
		self.admin = make_user(ADMIN)
		self.inv = make_user(INVESTIGADOR)
		self.analista = make_user(ANALISTA)
		Kazu.objects.create(munisipiu=Munisipiu.objects.get(code='LIQ'), status='VERIFIED', kode='RbR-LIQ-2026-0001',
			created_by=self.inv, konsentimentu=True, latitude='-8.6', longitude='125.3', urjente=True)

	def kria_link(self):
		self.client.force_login(self.admin)
		self.client.post(reverse('report-live-ekran-add'), {'naran': 'Sala reuniaun'})
		url = self.client.session.get('live_link_foun')['url']
		self.client.logout()
		return url.split('/live/')[1].strip('/')

	def test_link_kiosk_la_presiza_login(self):
		token = self.kria_link()
		r = self.client.get(reverse('live-kiosk', args=[token]))
		self.assertEqual(r.status_code, 200)
		self.assertEqual(r['Referrer-Policy'], 'no-referrer')
		d = self.client.get(reverse('live-kiosk-dadus', args=[token])).json()
		self.assertEqual(d['kpi']['total'], 1)
		self.assertEqual(d['kpi']['urjente'], 1)
		self.assertEqual(d['pontu'][0]['kode'], 'RbR-LIQ-2026-0001')
		# Dadus anónimu: la iha naran / email utilizador
		self.assertNotIn(self.inv.email, str(d))

	def test_token_sala_no_kansela(self):
		self.assertEqual(self.client.get(reverse('live-kiosk-dadus', args=['sala'])).status_code, 404)
		token = self.kria_link()
		from report.models import LiveEkran
		ek = LiveEkran.objects.get()
		self.assertNotEqual(ek.token_hash, token)          # token rai hanesan hash
		self.client.force_login(self.admin)
		self.client.post(reverse('report-live-ekran-kansela', args=[ek.pk]))
		self.client.logout()
		self.assertEqual(self.client.get(reverse('live-kiosk', args=[token])).status_code, 404)

	def test_live_staff_no_jestaun(self):
		self.assertEqual(self.client.get(reverse('report-live')).status_code, 302)
		self.client.force_login(self.analista)
		self.assertEqual(self.client.get(reverse('report-live')).status_code, 200)
		self.assertEqual(self.client.get(reverse('report-live-dadus')).status_code, 200)
		self.assertEqual(self.client.get(reverse('report-live-ekran')).status_code, 403)
		self.client.force_login(self.inv)
		self.assertEqual(self.client.get(reverse('report-live')).status_code, 403)
