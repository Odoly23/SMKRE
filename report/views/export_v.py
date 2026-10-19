import logging
from io import BytesIO
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from config.decorators import allowed_users
from config.rbac import ROLE_RELATORIU, ADMIN, SUPERADMIN
from report.utils import report_queryset
from users.auth_utils import c_user_group

logger = logging.getLogger('smkre.seguransa')


# Kontaktu pesoál: sai mamuk ba Analista / Ofisiál Legál (Admin no Superadmin haree hotu)
SUBAR_EXPORT = {'portavos_naran', 'portavos_kontaktu', 'portavos_observasaun', 'fokal_naran', 'fokal_telemovel',
	'org_seluk_naran', 'org_seluk_kontaktu', 'advogadu_kontaktu'}


@login_required
@allowed_users(allowed_roles=ROLE_RELATORIU)
def exportExcel(request):
	"""Eksporta kazu (tuir filtru) ho kolun hanesan baze dadus sentrál RbR (Airtable "Kazu Hotu") → bele import fali.
	Kolun SIBAR (ID, GPS, status…) iha ikus. Kontaktu pesoál: Admin / Superadmin de'it."""
	from kazu.airtable import KABESALLU, lina_export
	group = c_user_group(request.user)
	subar = set() if group in (ADMIN, SUPERADMIN) else SUBAR_EXPORT
	objects = report_queryset(request).select_related('munisipiu', 'postu', 'suku', 'aldeia', 'tipu_rai', 'created_by__pesoaluser__pesoal', 'segimentu') \
		.prefetch_related('tipu_konflitu', 'estragu', 'nesesidade', 'afetadu', 'insidente__tipu_eviksaun').order_by('referensia_rbr', 'kode')
	wb = Workbook()
	ws = wb.active
	ws.title = 'Kazu Hotu'
	# Kolun SIBAR: naran la hanesan kolun Airtable (import fali la sala)
	extra = [_('ID Kazu') + ' (SIBAR)', _('Latitude'), _('Longitude'), _('GPS aproksimadu'), _('Status Dadus') + ' (SIBAR)', _('Status Kazu') + ' (SIBAR)',
		_('Data Akontesimentu'), _('Tipu Rai'), _('Luan poligon (m²)'), _('Urjente'), _('Investigadór')]
	head = KABESALLU + [str(h) for h in extra]
	ws.append(head)
	verde, azul = PatternFill('solid', fgColor='0E8A68'), PatternFill('solid', fgColor='1F5F8B')
	for i, c in enumerate(ws[1]):
		c.font = Font(bold=True, color='FFFFFF')
		c.fill = verde if i < len(KABESALLU) else azul             # azul = kolun SIBAR
		c.alignment = Alignment(vertical='center', wrap_text=True)
	for k in objects:
		ws.append(lina_export(k, subar) + [
			k.kode or '', float(k.latitude) if k.latitude is not None else None, float(k.longitude) if k.longitude is not None else None,
			_('Sin') if k.gps_aproksimadu else _('Lae'), str(k.get_status_display()), str(k.get_status_kazu_display()),
			k.data_akontesimentu, k.tipu_rai.name if k.tipu_rai else '',
			float(k.poligon_luan_m2) if k.poligon_luan_m2 is not None else None, _('Sin') if k.urjente else _('Lae'),
			str(getattr(getattr(k.created_by, 'pesoaluser', None), 'pesoal', '') or k.created_by.username),
		])
	for i, _h in enumerate(head, start=1):
		ws.column_dimensions[get_column_letter(i)].width = 20
	ws.row_dimensions[1].height = 45
	ws.freeze_panes = 'B2'
	ws.auto_filter.ref = ws.dimensions
	buf = BytesIO()
	wb.save(buf)
	logger.info(f'Eksporta Excel: {objects.count()} kazu husi {request.user.username}')
	response = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
	response['Content-Disposition'] = f'attachment; filename="SIBAR-RbR-kazu-{timezone.localdate():%Y%m%d}.xlsx"'
	return response
