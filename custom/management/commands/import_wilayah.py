"""
Import Postu / Suku / Aldeia.

Formatu 1 — Diploma Ministerial 31/2026 (rekomenda):
	python manage.py import_wilayah custom/data/wilayah/wilayah_dm31_2026.xlsx --dry-run
	python manage.py import_wilayah custom/data/wilayah/wilayah_dm31_2026.xlsx
  ka pasta ho CSV: postu_dm31_2026.csv, suku_dm31_2026.csv, aldeia_dm31_2026.csv
	python manage.py import_wilayah custom/data/wilayah/

Formatu 2 — CSV ida (tuan): postu_code,suku_code,suku,aldeia_code,aldeia[,latitude,longitude]

Seguru atu la'o dala barak (update_or_create tuir kódigu). Suku ho postu_code mamuk (VERIFIKA) la tama.
"""
import csv
from pathlib import Path
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from custom.models import Munisipiu, PostuAdministrativu, Suku, Aldeia


def lee_csv(path):
	with open(path, encoding='utf-8-sig') as f:
		return list(csv.DictReader(f))


def lee_xlsx(path, aba):
	from openpyxl import load_workbook
	wb = load_workbook(path, read_only=True, data_only=True)
	if aba not in wb.sheetnames:
		raise CommandError(f'Aba "{aba}" la iha iha {path}')
	rows = list(wb[aba].iter_rows(values_only=True))
	kab = [str(c or '').strip() for c in rows[0]]
	return [{k: ('' if v is None else str(v).strip()) for k, v in zip(kab, r)} for r in rows[1:] if any(r)]


class Command(BaseCommand):
	help = "Import Postu, Suku no Aldeia (Excel DM 31/2026, pasta CSV, ka CSV tuan). Seguru atu la'o dala barak."

	def add_arguments(self, parser):
		parser.add_argument('fonte', help='Excel (.xlsx), pasta ho CSV DM 31/2026, ka CSV formatu tuan')
		parser.add_argument('--dry-run', action='store_true', help="Haree rezultadu deit, la rai buat ida")
		parser.add_argument('--desativa-la-iha', action='store_true',
			help="Desativa Suku/Aldeia iha base dadus ne'ebé la iha iha file (ex. dadus demo). Uza deit ho file kompletu.")

	def handle(self, *args, **opt):
		fonte = Path(opt['fonte'])
		if not fonte.exists():
			raise CommandError(f'File ka pasta la iha: {fonte}')
		if fonte.is_dir():
			base = {n: next(iter(sorted(fonte.glob(f'{n}_*.csv'))), None) for n in ('postu', 'suku', 'aldeia')}
			if not base['suku']:
				raise CommandError(f'La hetan suku_*.csv iha {fonte}')
			dados = {n: (lee_csv(p) if p else []) for n, p in base.items()}
		elif fonte.suffix.lower() == '.xlsx':
			dados = {'postu': lee_xlsx(fonte, 'Postu'), 'suku': lee_xlsx(fonte, 'Suku'), 'aldeia': lee_xlsx(fonte, 'Aldeia')}
		else:
			return self.formatu_tuan(fonte, opt['dry_run'])

		with transaction.atomic():
			res = self.importa(dados, opt['desativa_la_iha'])
			if opt['dry_run']:
				transaction.set_rollback(True)
		self.relatoriu(res, opt['dry_run'])

	# ── Formatu DM 31/2026 ──
	def importa(self, d, desativa):
		r = {'postu_foun': 0, 'suku_foun': 0, 'suku_atualiza': 0, 'aldeia_foun': 0, 'aldeia_atualiza': 0,
			'verifika': [], 'erru': [], 'desativa_suku': 0, 'desativa_aldeia': 0}
		mun = {m.code: m for m in Munisipiu.objects.all()}
		for p in d['postu']:
			m = mun.get(p['munisipiu'])
			if not m:
				r['erru'].append(f'Postu {p["code"]}: munisípiu {p["munisipiu"]} la iha (halo setup_smkre uluk)')
				continue
			_, foun = PostuAdministrativu.objects.update_or_create(code=p['code'], defaults={'name': p['name'], 'munisipiu': m})
			r['postu_foun'] += int(foun)
		postu = {p.code: p for p in PostuAdministrativu.objects.select_related('munisipiu')}

		suku_ok = {}
		for s in d['suku']:
			code = s['suku_code']
			if not s.get('postu_code'):
				r['verifika'].append(f"{code} {s['suku']} ({s.get('munisipiu_naran') or s.get('munisipiu')})")
				continue
			p = postu.get(s['postu_code'])
			if not p:
				r['erru'].append(f'{code} {s["suku"]}: postu_code "{s["postu_code"]}" la iha')
				continue
			if s.get('munisipiu') and p.munisipiu.code != s['munisipiu']:
				r['erru'].append(f'{code} {s["suku"]}: postu {p.code} la iha munisípiu {s["munisipiu"]}')
				continue
			obj, foun = Suku.objects.update_or_create(code=code, defaults={
				'name': s['suku'], 'postu': p, 'order': int(float(s.get('no') or 0)), 'is_active': True})
			r['suku_foun' if foun else 'suku_atualiza'] += 1
			suku_ok[code] = obj

		for a in d['aldeia']:
			s = suku_ok.get(a['suku_code'])
			if not s:
				continue                                  # suku seidauk tama (VERIFIKA)
			_, foun = Aldeia.objects.update_or_create(code=a['aldeia_code'], defaults={
				'name': a['aldeia'], 'suku': s, 'order': int(float(a.get('no') or 0)), 'is_active': True})
			r['aldeia_foun' if foun else 'aldeia_atualiza'] += 1

		if desativa:
			if r['verifika']:
				raise_msg = 'La bele uza --desativa-la-iha bainhira iha suku VERIFIKA seidauk prenxe.'
				r['erru'].append(raise_msg)
			else:
				codes_s = {s['suku_code'] for s in d['suku']}
				codes_a = {a['aldeia_code'] for a in d['aldeia']}
				r['desativa_suku'] = Suku.objects.filter(is_active=True).exclude(code__in=codes_s).update(is_active=False)
				r['desativa_aldeia'] = Aldeia.objects.filter(is_active=True).exclude(code__in=codes_a).update(is_active=False)
		return r

	def relatoriu(self, r, dry):
		for e in r['erru']:
			self.stderr.write(self.style.ERROR(f'ERRU: {e}'))
		if r['verifika']:
			self.stdout.write(self.style.WARNING(f'{len(r["verifika"])} suku seidauk iha postu (VERIFIKA) — la tama:'))
			for v in r['verifika'][:15]:
				self.stdout.write(f'   · {v}')
			if len(r['verifika']) > 15:
				self.stdout.write(f'   … no {len(r["verifika"]) - 15} tan. Prenxe postu_code iha aba "Suku" no import fali.')
		self.stdout.write(self.style.SUCCESS(
			f"{'[DRY-RUN — la rai buat ida] ' if dry else ''}"
			f"Postu foun: {r['postu_foun']} · Suku foun: {r['suku_foun']} (atualiza {r['suku_atualiza']}) · "
			f"Aldeia foun: {r['aldeia_foun']} (atualiza {r['aldeia_atualiza']})"
			+ (f" · Desativa: suku {r['desativa_suku']}, aldeia {r['desativa_aldeia']}" if r['desativa_suku'] or r['desativa_aldeia'] else '')))

	# ── Formatu tuan (CSV ida) ──
	def formatu_tuan(self, path, dry):
		n_suku, n_aldeia, erru = 0, 0, []
		with transaction.atomic():
			for line, row in enumerate(lee_csv(path), start=2):
				try:
					postu = PostuAdministrativu.objects.get(code=row['postu_code'].strip())
				except (PostuAdministrativu.DoesNotExist, KeyError):
					erru.append(f'liña {line}: postu "{row.get("postu_code")}" la iha')
					continue
				suku, created = Suku.objects.update_or_create(code=row['suku_code'].strip(), defaults={
					'name': row['suku'].strip(), 'postu': postu,
					'latitude': (row.get('latitude') or '').strip() or None,
					'longitude': (row.get('longitude') or '').strip() or None})
				n_suku += int(created)
				if (row.get('aldeia_code') or '').strip():
					_, created = Aldeia.objects.update_or_create(code=row['aldeia_code'].strip(), defaults={
						'name': row['aldeia'].strip(), 'suku': suku})
					n_aldeia += int(created)
			if dry:
				transaction.set_rollback(True)
		for e in erru:
			self.stderr.write(self.style.WARNING(e))
		self.stdout.write(self.style.SUCCESS(f"{'[DRY-RUN] ' if dry else ''}Suku foun: {n_suku} · Aldeia foun: {n_aldeia} · Erru: {len(erru)}"))
