import tempfile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from config.rbac import INVESTIGADOR
from config.testing import setup_master, make_user
from custom.models import Munisipiu, PostuAdministrativu, Suku, Aldeia, TipuKonflitu


class SetupTest(TestCase):
	def test_setup_kria_dadus_no_bele_la_o_dala_rua(self):
		setup_master()
		setup_master()
		self.assertEqual(Munisipiu.objects.count(), 14)
		self.assertEqual(PostuAdministrativu.objects.count(), 65)
		self.assertTrue(TipuKonflitu.objects.get(code='SELUK').presiza_esplika)


class WilayahTest(TestCase):
	def setUp(self):
		setup_master()
		with tempfile.NamedTemporaryFile('w', suffix='.csv', delete=False, encoding='utf-8') as f:
			f.write('postu_code,suku_code,suku,aldeia_code,aldeia\n')
			f.write('LIQ-03,LIQ-03-01,Vatuvou,LIQ-03-01-01,Aldeia A\n')
			f.write('LIQ-03,LIQ-03-01,Vatuvou,LIQ-03-01-02,Aldeia B\n')
			f.write('XXX-99,XXX,Sala,,\n')
			self.csv = f.name
		call_command('import_wilayah', self.csv, stdout=open('/dev/null', 'w'), stderr=open('/dev/null', 'w'))

	def test_import_suku_aldeia(self):
		self.assertEqual(Suku.objects.count(), 1)
		self.assertEqual(Aldeia.objects.count(), 2)

	def test_dropdown_presiza_login(self):
		r = self.client.get(reverse('ajax_load_suku'), {'postu': 1})
		self.assertEqual(r.status_code, 302)

	def test_dropdown_bertingkat(self):
		self.client.force_login(make_user(INVESTIGADOR))
		postu = PostuAdministrativu.objects.get(code='LIQ-03')
		r = self.client.get(reverse('ajax_load_suku'), {'postu': postu.pk})
		self.assertContains(r, 'Vatuvou')
		r = self.client.get(reverse('ajax_load_post'), {'munisipiu': postu.munisipiu_id})
		self.assertContains(r, 'Maubara')
		r = self.client.get(reverse('ajax_load_aldeia'), {'suku': Suku.objects.get().pk})
		self.assertContains(r, 'Aldeia B')


class ImportDM31Test(TestCase):
	"""Import wilayah Diploma Ministerial 31/2026 (custom/data/wilayah/)."""
	XLSX = 'custom/data/wilayah/wilayah_dm31_2026.xlsx'

	def setUp(self):
		setup_master()

	def run_cmd(self, *args):
		from io import StringIO
		out = StringIO()
		call_command('import_wilayah', *args, stdout=out, stderr=StringIO())
		return out.getvalue()

	def test_dry_run_la_rai(self):
		self.run_cmd(self.XLSX, '--dry-run')
		self.assertEqual(Suku.objects.count(), 0)
		self.assertEqual(PostuAdministrativu.objects.count(), 65)

	def test_import_idempotente(self):
		out = self.run_cmd(self.XLSX)
		self.assertIn('Postu foun: 6', out)
		self.assertEqual(Suku.objects.count(), 472)                 # VERIFIKA hotu prenxe ona (fonte web)
		self.assertEqual(Aldeia.objects.count(), 2250)
		self.assertEqual(PostuAdministrativu.objects.get(code='BAU-07').name, 'Quelicai Antigu')
		self.assertEqual(Suku.objects.get(code='LIQ-S344').name, 'Maubaralissa')   # PDF hakerek 354
		self.assertEqual(Suku.objects.get(code='LAU-S299').postu.code, 'LAU-02')    # Baduro → Lautém
		self.assertEqual(Suku.objects.get(code='DIL-S230').postu.code, 'DIL-02')    # Madohi → Dom Aleixo
		out = self.run_cmd(self.XLSX)                               # la'o fali → la duplika
		self.assertIn('Suku foun: 0', out)
		self.assertEqual(Suku.objects.count(), 472)

	def test_verifika_prenxe_tuir_diploma_no_fonte(self):
		# Postu ne'ebé prenxe tenke iha opsaun diploma (sujestaun) no obs tenke iha fonte
		import csv
		import re
		with open('custom/data/wilayah/suku_dm31_2026.csv', encoding='utf-8') as f:
			rows = list(csv.DictReader(f))
		self.assertEqual(len(rows), 472)
		self.assertFalse([r['suku_code'] for r in rows if not r['postu_code']])
		prenxe = [r for r in rows if r['sujestaun']]
		self.assertEqual(len(prenxe), 57)
		for r in prenxe:
			self.assertIn(r['verifika'], ('WEB', 'DIPLOMA'), r['suku_code'])
			self.assertIn(r['postu_code'], re.findall(r'[A-Z]{3}-\d\d', r['sujestaun']), r['suku_code'])
			self.assertIn('Fonte postu:', r['obs'], r['suku_code'])
			if r['verifika'] == 'WEB':
				self.assertRegex(r['obs'], r'Fonte postu: https://', r['suku_code'])

	def test_desativa_demo(self):
		demo = Suku.objects.create(code='DEMO-01', name='Demo', postu=PostuAdministrativu.objects.first())
		self.run_cmd(self.XLSX, '--desativa-la-iha')
		self.assertEqual(Suku.objects.filter(is_active=True).count(), 472)
		self.assertEqual(Aldeia.objects.filter(is_active=True).count(), 2250)
		demo.refresh_from_db()
		self.assertFalse(demo.is_active)

	def test_desativa_la_bele_ho_verifika(self):
		# Se iha suku VERIFIKA mamuk, --desativa-la-iha la desativa buat ida
		from openpyxl import load_workbook
		wb = load_workbook(self.XLSX)
		ws = wb['Suku']
		ip = [c.value for c in ws[1]].index('postu_code')
		ws.cell(row=2, column=ip + 1).value = None
		tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
		wb.save(tmp.name)
		demo = Suku.objects.create(code='DEMO-01', name='Demo', postu=PostuAdministrativu.objects.first())
		self.run_cmd(tmp.name, '--desativa-la-iha')
		self.assertEqual(Suku.objects.filter(is_active=True).count(), 472)   # 471 + demo
		demo.refresh_from_db()
		self.assertTrue(demo.is_active)
