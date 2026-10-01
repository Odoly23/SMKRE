"""Prenxe postu_code ba suku VERIFIKA tuir fonte web (Outubru 2026) → atualiza CSV + Excel.

Diploma Ministerial 31/2026 mak referénsia prinsipál: postu ne'ebé hili tenke mós iha opsaun husi diploma
(kolun `sujestaun`). Se fonte web kontra diploma, la muda — hakerek iha `obs` deit.

Nota: situs munisípiu *.gov.tl (lautem/ermera/aileu/…) la bele asesu husi ambiente ne'ebé halo atualizasaun ne'e
(network bloqueadu); fonte sira hetan liu pesquiza web (Wikipedia, citypopulation.de — dadus Censo 2015/2022, ka notísia).

	python tools/wilayah_dm31/5_fonte_web.py custom/data/wilayah
"""
import csv
import re
import sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

OUT = sys.argv[1] if len(sys.argv) > 1 else 'custom/data/wilayah'

W = 'https://en.wikipedia.org/wiki/'
CP = 'https://www.citypopulation.de/en/timor/admin/'
REMEXIO = f'{W}Remexio_Administrative_Post ; {CP}0104__remexio'
LEQUIDOE = f'{W}Lequidoe_Administrative_Post ; {CP}0103__lequidoe'
LAULARA = f'{W}Laulara_Administrative_Post ; {CP}0102__laulara/'
LAUTEM = f'{W}Laut%C3%A9m_Administrative_Post'
NAINFETO = f'{W}Nain_Feto_Administrative_Post ; {CP}0605__nain_feto/'
BOBONARO = f'{W}Bobonaro_Administrative_Post ; {CP}0403__bobonaro'
CAILACO = f'{W}Cailaco_Administrative_Post'
ERMERA = f'{W}Ermera_Administrative_Post'
RAILACO = f'{W}Railaco_Administrative_Post'
TURISCAI = f'{W}Turiscai_Administrative_Post'

# suku_code: (postu_code, fonte, nota)
FONTE = {
	'AIL-S015': ('AIL-04', REMEXIO, ''),
	'AIL-S016': ('AIL-04', 'https://de.wikipedia.org/wiki/Tulataqueo', 'Aicurus uluk aldeia husi suco Tulataqueo (Remexio)'),
	'AIL-S017': ('AIL-04', 'https://timorpost.com/jeral-tl/tp-54977/pnds-postu-remexio-ho-ks-estabelese-es-pnds-iha-suku-karahili/', 'Suku Karahili iha Postu Remexio'),
	'AIL-S018': ('AIL-04', REMEXIO, ''),
	'AIL-S019': ('AIL-04', REMEXIO, ''),
	'AIL-S020': ('AIL-04', REMEXIO, ''),
	'AIL-S021': ('AIL-04', REMEXIO, 'Fonte web hakerek "Faturasa"'),
	'AIL-S022': ('AIL-04', REMEXIO, ''),
	'AIL-S023': ('AIL-04', REMEXIO, ''),
	'AIL-S024': ('AIL-04', REMEXIO, ''),
	**{f'AIL-S{n:03d}': ('AIL-03', LEQUIDOE, '') for n in range(25, 32)},
	**{f'AIL-S{n:03d}': ('AIL-02', LAULARA, '') for n in range(32, 38)},
	'BAU-S095': ('BAU-03', f'{W}Laga_Administrative_Post ; https://citypopulation.de/en/timor/admin/0303__laga/', ''),
	'BOB-S146': ('BOB-03', BOBONARO, ''),
	'BOB-S147': ('BOB-03', BOBONARO, ''),
	'BOB-S155': ('BOB-04', CAILACO, ''),
	'BOB-S162': ('BOB-04', CAILACO, ''),
	'DIL-S216': ('DIL-04', NAINFETO, ''),
	'DIL-S217': ('DIL-04', NAINFETO, ''),
	'DIL-S218': ('DIL-04', NAINFETO, ''),
	'DIL-S221': ('DIL-04', NAINFETO, ''),
	'DIL-S230': ('DIL-02', f'{W}Dom_Aleixo_Administrative_Post ; {CP}0603__dom_aleixo/', ''),
	'ERM-S247': ('ERM-02', ERMERA, 'Fonte web hakerek suco ida "Poetete" (Postu Ermera)'),
	'ERM-S248': ('ERM-02', ERMERA, 'Fonte web hakerek suco ida "Poetete" (Postu Ermera)'),
	'ERM-S275': ('ERM-05', RAILACO, ''),
	'ERM-S283': ('ERM-05', RAILACO, ''),
	**{f'LAU-S{n:03d}': ('LAU-02', LAUTEM, '') for n in range(299, 309)},
	'LAU-S309': ('LAU-03', f'{W}Lospalos_Administrative_Post ; {CP}0803__lospalos', ''),
	'LIQ-S336': ('LIQ-01', f'{W}Bazartete_Administrative_Post', 'Fonte web hakerek "Fahilebo"'),
	'LIQ-S337': ('LIQ-02', f'{W}Liqui%C3%A7%C3%A1_Administrative_Post', ''),
	'MAN-S365': ('MAN-02', 'https://citypopulation.de/en/timor/admin/laclo/100205__uma_naruc/', ''),
	'MAN-S372': ('MAN-06', 'https://de.wikipedia.org/wiki/Fatumaquerec_(Soibada)', 'Aldeia Lesuata no Sasahi hanesan ho fonte'),
	'MNF-S392': ('MNF-03', '', 'La hetan fonte web (suco foun). Postu tuir orden alfabétiku iha anexu diploma: '
		'Same (Babulo…Tutuluro, Uetano) — Fatuberlio hahú ho Bubususo'),
	'MNF-S398': ('MNF-04', TURISCAI, ''),
	'MNF-S408': ('MNF-04', TURISCAI, ''),
	'VIQ-S454': ('VIQ-02', 'https://wikipedia.jakami.de/content/wikipedia_de_all_mini_2024-06/A/Uaibobo', ''),
	'VIQ-S463': ('VIQ-03', f'{W}Uatucarbau_Administrative_Post', 'Afaloicai seluk (VIQ-S455) iha Uatolari'),
}


def opsaun_diploma(sujestaun):
	return set(re.findall(r'[A-Z]{3}-\d\d', sujestaun or ''))


def main():
	postu = {r['code']: r['name'] for r in csv.DictReader(open(f'{OUT}/postu_dm31_2026.csv', encoding='utf-8'))}
	path = f'{OUT}/suku_dm31_2026.csv'
	rows = list(csv.DictReader(open(path, encoding='utf-8')))
	cols = list(rows[0].keys())
	muda = {}
	for r in rows:
		if r['verifika'] != 'SIN':
			continue
		if r['suku_code'] not in FONTE:
			print(f"La iha fonte: {r['suku_code']} {r['suku']}")
			continue
		code, url, nota = FONTE[r['suku_code']]
		obs = [r['obs']] if r['obs'] else []
		if code not in opsaun_diploma(r['sujestaun']):
			obs.append(f'Fonte web hatudu {code} maibé diploma la fó opsaun ne\'e — la muda (diploma mak referénsia): {url}')
			r['obs'] = '; '.join(obs)
			muda[r['suku_code']] = r
			continue
		if nota:
			obs.append(nota)
		obs.append(f'Fonte postu: {url}' if url else 'Fonte postu: Diploma Ministerial 31/2026 (orden anexu)')
		r.update(postu_code=code, postu=postu[code], verifika='WEB' if url else 'DIPLOMA', obs='; '.join(obs))
		muda[r['suku_code']] = r
	with open(path, 'w', encoding='utf-8', newline='') as f:
		w = csv.DictWriter(f, fieldnames=cols)
		w.writeheader()
		w.writerows(rows)

	xlsx = f'{OUT}/wilayah_dm31_2026.xlsx'
	wb = load_workbook(xlsx)
	ws = wb['Suku']
	kab = [c.value for c in ws[1]]
	ik = kab.index('suku_code')
	MEAR = PatternFill('solid', fgColor='DFF5E8')
	for row in ws.iter_rows(min_row=2):
		r = muda.get(row[ik].value)
		if not r:
			continue
		for c, k in zip(row, kab):
			c.value = r[k] if k not in ('no', 'n_aldeia') else int(r[k])
		if r['verifika'] != 'SIN':
			for c in row:
				c.fill = MEAR
	resta = sum(r['verifika'] == 'SIN' for r in rows)
	if not muda:
		print(f'La iha buat ida atu muda · VERIFIKA resta {resta}')
		return
	info = wb['Lee-uluk']
	info['A3'] = re.sub(r'Suku VERIFIKA: \d+.*', f'Suku VERIFIKA: {resta} (prenxe ona {len(muda)} tuir fonte web — haree kolun obs)', info['A3'].value)
	info['A5'] = ('1. Aba "Suku": liña kór-mean (verifika = SIN) → hakerek postu_code. Liña kór-verde (verifika = WEB/DIPLOMA) '
		'prenxe ona tuir fonte iha kolun obs.')
	wb.save(xlsx)
	print(f'Prenxe {len(muda)} suku · VERIFIKA resta {resta}')


if __name__ == '__main__':
	main()
